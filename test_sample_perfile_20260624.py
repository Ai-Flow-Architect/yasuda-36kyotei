"""
追加開発「事業所ごとの見本出し分け」回帰テスト（2026-06-24）

スコープ:
  1) sample_selector のトークン分解・マッチング・選別
  2) excel_reader が「見本指定」列を位置非依存で読み取る
  3) 後方互換: 見本指定が無ければ従来の全宛先共通モード判定
"""
import tempfile
from pathlib import Path

import openpyxl

from excel_reader import read_excel, _detect_sample_spec_column, SAMPLE_SPEC_COLUMN
from sample_selector import (
    parse_sample_spec,
    select_samples_for_record,
    has_any_sample_spec,
    build_extra_attachments,
)


# プール（bytes, filename, mime）のダミー
def _pool(*names: str) -> list[tuple[bytes, str, str]]:
    return [(b"dummy", n, "application/pdf") for n in names]


# ============================================================
# 1) sample_selector
# ============================================================
def test_parse_spec_separators():
    """各種区切りでトークン分解できる"""
    assert parse_sample_spec("a.pdf、b.pdf") == ["a.pdf", "b.pdf"]
    assert parse_sample_spec("a.pdf, b.pdf; c.pdf") == ["a.pdf", "b.pdf", "c.pdf"]
    assert parse_sample_spec("a.pdf / b.pdf") == ["a.pdf", "b.pdf"]
    assert parse_sample_spec("a.pdf\nb.pdf") == ["a.pdf", "b.pdf"]


def test_parse_spec_empty_and_dedupe():
    """空・重複の正規化"""
    assert parse_sample_spec("") == []
    assert parse_sample_spec("   ") == []
    assert parse_sample_spec(None) == []
    assert parse_sample_spec("a.pdf、a.pdf") == ["a.pdf"]


def test_select_exact_filename():
    """完全ファイル名一致"""
    pool = _pool("建設業見本.pdf", "製造業見本.pdf")
    sel, un = select_samples_for_record("建設業見本.pdf", pool)
    assert [s[1] for s in sel] == ["建設業見本.pdf"]
    assert un == []


def test_select_stem_match():
    """拡張子なし指定でも一致（建設業見本 ⇔ 建設業見本.pdf）"""
    pool = _pool("建設業見本.pdf", "製造業見本.pdf")
    sel, un = select_samples_for_record("建設業見本", pool)
    assert [s[1] for s in sel] == ["建設業見本.pdf"]
    assert un == []


def test_select_multiple():
    """複数指定で複数添付（pool順を保持）"""
    pool = _pool("建設業見本.pdf", "製造業見本.pdf", "運送業見本.pdf")
    sel, un = select_samples_for_record("製造業見本、運送業見本", pool)
    assert [s[1] for s in sel] == ["製造業見本.pdf", "運送業見本.pdf"]
    assert un == []


def test_select_partial_match():
    """部分一致（指定がファイル名stemに含まれる）"""
    pool = _pool("2026_建設業_見本一式.pdf")
    sel, un = select_samples_for_record("建設業", pool)
    assert len(sel) == 1
    assert un == []


def test_select_unmatched_reported():
    """プールに無い指定は未一致として報告（サイレント欠落防止）"""
    pool = _pool("建設業見本.pdf")
    sel, un = select_samples_for_record("存在しない見本.pdf", pool)
    assert sel == []
    assert un == ["存在しない見本.pdf"]


def test_select_zenkaku_hankaku_case():
    """全角半角・大文字小文字の揺れを吸収（NFKC）"""
    pool = _pool("SampleA.pdf")
    sel, un = select_samples_for_record("ｓａｍｐｌｅａ", pool)  # 全角小文字
    assert len(sel) == 1
    assert un == []


def test_select_empty_spec_no_attach():
    """見本指定が空なら何も添付しない（出し分けモードの安全側）"""
    pool = _pool("a.pdf")
    sel, un = select_samples_for_record("", pool)
    assert sel == []
    assert un == []


def test_select_no_duplicate_when_two_tokens_hit_same_file():
    """2トークンが同一ファイルに当たっても重複添付しない"""
    pool = _pool("建設業見本.pdf")
    sel, un = select_samples_for_record("建設業見本.pdf、建設業", pool)
    assert [s[1] for s in sel] == ["建設業見本.pdf"]


def test_select_exact_beats_partial_no_overattach():
    """完全/拡張子なし一致があれば、ゆるい部分一致のファイルは巻き込まない（過剰添付防止・3AI共通指摘の底上げ）"""
    pool = _pool("見本.pdf", "建設業見本.pdf", "製造業見本.pdf")
    # 「見本.pdf」は1件目に完全一致。部分一致では2,3件目も該当しうるが採用しない。
    sel, un = select_samples_for_record("見本.pdf", pool)
    assert [s[1] for s in sel] == ["見本.pdf"]
    assert un == []


def test_select_partial_only_when_no_exact():
    """完全/stem一致が無いトークンは部分一致で複数拾う（従来の利便は維持）"""
    pool = _pool("2026建設業見本.pdf", "建設業安全見本.pdf", "製造業見本.pdf")
    sel, un = select_samples_for_record("建設業", pool)
    assert {s[1] for s in sel} == {"2026建設業見本.pdf", "建設業安全見本.pdf"}
    assert un == []


def test_normalize_none_safe():
    """None指定が 'None' 文字列として誤マッチしない"""
    pool = _pool("None.pdf", "見本.pdf")
    sel, un = select_samples_for_record(None, pool)
    assert sel == [] and un == []


def test_has_any_sample_spec():
    """1件でも見本指定があればTrue（モード判定）"""
    assert has_any_sample_spec([{"見本指定": ""}, {"見本指定": "a.pdf"}]) is True
    assert has_any_sample_spec([{"見本指定": ""}, {}]) is False
    assert has_any_sample_spec([]) is False


# ============================================================
# 2) excel_reader: 見本指定列の読み取り（位置非依存）
# ============================================================
def _make_excel_with_sample_col(path: str, *, header: str | None, col: int, value: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=4, value="事業所名")
    ws.cell(row=1, column=5, value="事業主名")
    ws.cell(row=1, column=44, value="送信先メールアドレス")
    if header:
        ws.cell(row=1, column=col, value=header)
    ws.cell(row=2, column=4, value="株式会社サンプル商事")
    ws.cell(row=2, column=5, value="田中一郎")
    ws.cell(row=2, column=44, value="to@example.invalid")
    if header:
        ws.cell(row=2, column=col, value=value)
    wb.save(path)


def test_excel_sample_spec_detected_by_header():
    """『見本指定』ヘッダーを位置非依存で検出して読み取る（I列に配置）"""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as t:
        path = t.name
    _make_excel_with_sample_col(path, header="見本指定", col=9, value="建設業見本.pdf")
    records, _ = read_excel(path)
    assert records[0]["見本指定"] == "建設業見本.pdf"


def test_excel_sample_spec_header_variant():
    """表記ゆれ『添付見本』も検出する"""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as t:
        path = t.name
    _make_excel_with_sample_col(path, header="添付見本", col=10, value="製造業見本.pdf")
    records, _ = read_excel(path)
    assert records[0]["見本指定"] == "製造業見本.pdf"


def test_excel_no_sample_col_backward_compat():
    """見本指定列が無ければ空文字（従来Excelで後方互換・共通モード）"""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as t:
        path = t.name
    _make_excel_with_sample_col(path, header=None, col=0, value="")
    records, _ = read_excel(path)
    assert records[0]["見本指定"] == ""
    assert has_any_sample_spec(records) is False


def test_excel_sample_spec_fixed_au_column():
    """ヘッダー未検出でもAU列(47)固定フォールバックで読む"""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as t:
        path = t.name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=4, value="事業所名")
    ws.cell(row=2, column=4, value="株式会社サンプル商事")
    # ヘッダー名を付けずにAU列(47)へ値だけ入れる
    ws.cell(row=2, column=SAMPLE_SPEC_COLUMN, value="運送業見本.pdf")
    wb.save(path)
    records, _ = read_excel(path)
    assert records[0]["見本指定"] == "運送業見本.pdf"


# ============================================================
# 3) build_extra_attachments（添付組み立ての純ロジック）
# ============================================================
def _kyotei(*names):
    return [(b"k", n, "application/pdf") for n in names]


def test_attach_common_mode_uses_all_samples():
    """共通モード: 全見本＋協定書2件目以降を結合（従来動作）"""
    pool = _pool("見本A.pdf", "見本B.pdf")
    extra, selected, un = build_extra_attachments(
        _kyotei("協定書2.pdf"), pool, sample_spec="", per_office=False
    )
    assert [a[1] for a in extra] == ["協定書2.pdf", "見本A.pdf", "見本B.pdf"]
    assert [s[1] for s in selected] == ["見本A.pdf", "見本B.pdf"]
    assert un == []


def test_attach_per_office_mode_filters_samples():
    """出し分けモード: 指定見本のみ＋協定書2件目以降"""
    pool = _pool("見本A.pdf", "見本B.pdf")
    extra, selected, un = build_extra_attachments(
        _kyotei("協定書2.pdf"), pool, sample_spec="見本B.pdf", per_office=True
    )
    assert [a[1] for a in extra] == ["協定書2.pdf", "見本B.pdf"]
    assert [s[1] for s in selected] == ["見本B.pdf"]
    assert un == []


def test_attach_per_office_unmatched_reported_and_excluded():
    """出し分けで未一致指定は添付されず報告される"""
    pool = _pool("見本A.pdf")
    extra, selected, un = build_extra_attachments(
        [], pool, sample_spec="存在しない.pdf", per_office=True
    )
    assert extra == []
    assert selected == []
    assert un == ["存在しない.pdf"]


def test_attach_per_office_empty_spec_no_sample():
    """出し分けで指定空なら見本は付かない（協定書のみ）"""
    pool = _pool("見本A.pdf")
    extra, selected, un = build_extra_attachments(
        _kyotei("協定書1.pdf"), pool, sample_spec="", per_office=True
    )
    assert [a[1] for a in extra] == ["協定書1.pdf"]
    assert selected == []


# ============================================================
# 4) デモExcel統合（生成物が出し分けで正しく解決する）
# ============================================================
def test_demo_excel_resolves_per_office():
    """make_perfile_demo の生成物を読み、指定どおりの見本が選ばれる"""
    import subprocess
    import sys
    base = Path(__file__).parent
    demo_xlsx = base / "demo_data" / "demo_perfile_36kyotei.xlsx"
    samples_dir = base / "demo_data" / "見本サンプル"
    # 未生成なら生成（再実行可能・冪等）
    if not demo_xlsx.exists() or not samples_dir.exists():
        subprocess.run([sys.executable, str(base / "make_perfile_demo.py")], check=True, cwd=str(base))

    records, _ = read_excel(str(demo_xlsx))
    assert has_any_sample_spec(records) is True
    pool = [(b"x", p.name, "application/pdf") for p in sorted(samples_dir.glob("*.pdf"))]

    by_name = {r["事業所名"]: r for r in records}
    sel1, un1 = select_samples_for_record(by_name["建設工業株式会社"]["見本指定"], pool)
    assert [s[1] for s in sel1] == ["建設業見本.pdf"] and un1 == []
    sel2, un2 = select_samples_for_record(by_name["テスト製作所"]["見本指定"], pool)
    assert [s[1] for s in sel2] == ["製造業見本.pdf"] and un2 == []
    sel3, _ = select_samples_for_record(by_name["ABCサービス合同会社"]["見本指定"], pool)
    assert {s[1] for s in sel3} == {"サービス業見本.pdf", "建設業見本.pdf"}


if __name__ == "__main__":
    import sys
    import pytest
    sys.exit(pytest.main([__file__, "-q"]))
