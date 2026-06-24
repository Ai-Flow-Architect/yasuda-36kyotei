"""
make_perfile_demo.py — 事業所別 見本出し分け の「テンプレート＋デモ一式」生成

受注確定後、客様がそのまま使えるよう以下を生成する（再実行で再生成可）:
  1. templates/36協定管理_見本出し分け対応_テンプレート.xlsx
       … 見本指定列を備えた空テンプレート（記入例コメント付き）
  2. demo_data/demo_perfile_36kyotei.xlsx
       … 3事業所に異なる見本を指定したデモ用Excel
  3. demo_data/見本サンプル/{建設業見本,製造業見本,サービス業見本}.pdf
       … 出し分けを試すためのダミー見本PDF（中身は最小の有効PDF）

固定列の整合（excel_reader.py と一致させること）:
  44=送信先メールアドレス / 45=様式パターン上書き / 46=事業所番号 / 47=見本指定
担当者名は本デモでは事業主名(E列)が人名のためフォールバックで宛名になる。
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

BASE = Path(__file__).parent
TEMPLATES = BASE / "templates"
DEMO = BASE / "demo_data"
SAMPLES = DEMO / "見本サンプル"

# excel_reader.py のCOLUMN_MAP（A〜AQ=43列）に対応する見出し + 拡張列
HEADERS = [
    "更新月", "起算日/月", "起算日/年", "事業所名", "事業主名",
    "電話番号", "事業の種類",
    "時間外労働をさせる必要のある具体的事由", "業務の種類①", "労働者数",
    "18歳未満の労働者数", "所定労働時間",
    "延長することができる時間/1日", "延長することができる時間/1ヶ月", "期間",
    "休日労働をさせる必要のある具体的事由", "業務の種類②",
    "所定休日", "休日労働させることができる日数", "始業及び終業時間", "期間",
    "特別条項の有無",
    "臨時的に限度時間を超える理由", "業務の種類", "労働者数",
    "延長できる時間", "限度時間を超えることができる回数",
    "延長することができる時間数/月", "限度時間を超えた場合の割増賃金率",
    "延長することができる時間数/年", "特別条項の手続き",
    "健康措置/該当番号", "健康措置/具体的内容",
    "誓約チェック",
    "労働者代表者職", "労働者代表者　氏名", "過半数労働者専任チェック",
    "協定締結日", "届出作成日",
    "事業主職名", "事業主名", "所轄労働局", "所轄労働基準監督署名",
    "送信先メールアドレス",     # 44 AR
    "様式パターン上書き",        # 45 AS（任意）
    "事業所番号",                # 46 AT（任意・ファイル名先頭番号と照合）
    "見本指定",                  # 47 AU（事業所ごとに添付する見本名）
]

# 最小限の有効な1ページPDF（xrefオフセットをコード側で正確に計算して組み立てる）
def _minimal_pdf(title: str) -> bytes:
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        None,  # 4: contents（後で本文差し込み）
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = b"BT /F1 14 Tf 30 110 Td (SAMPLE: " + title.encode("ascii", "replace") + b") Tj ET"
    objs[3] = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += (b"trailer\n<< /Size " + str(len(objs) + 1).encode()
            + b" /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF")
    return bytes(out)


def _write_headers(ws):
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["AU"].width = 22


def build_template():
    TEMPLATES.mkdir(exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "管理シート"
    _write_headers(ws)
    # 記入例（1行）— 見本指定の書き方を示す
    ws.cell(row=2, column=4, value="（例）建設工業株式会社")
    ws.cell(row=2, column=5, value="（例）サンプル太郎")
    ws.cell(row=2, column=44, value="（例）info@example.co.jp")
    ws.cell(row=2, column=46, value="0001")
    ws.cell(row=2, column=47, value="建設業見本.pdf")
    path = TEMPLATES / "36協定管理_見本出し分け対応_テンプレート.xlsx"
    wb.save(path)
    return path


def build_demo():
    DEMO.mkdir(exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "管理シート"
    _write_headers(ws)
    rows = [
        # (事業所名, 事業主名, 事業種類, 業務種類, 人数, メール, 事業所番号, 見本指定)
        ("建設工業株式会社", "サンプル太郎", "建設業", "現場管理", "20",
         "demo1@kensetsu.example.invalid", "0001", "建設業見本.pdf"),
        ("テスト製作所", "サンプル次郎", "製造業", "製造・品質管理", "25",
         "demo2@seizo.example.invalid", "0002", "製造業見本"),
        ("ABCサービス合同会社", "サンプル花子", "サービス業", "カスタマーサポート", "8",
         "demo3@service.example.invalid", "0003", "サービス業見本.pdf、建設業見本.pdf"),
    ]
    for i, (name, owner, gyoshu, gyomu, nin, mail, num, spec) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value="4月")
        ws.cell(row=i, column=4, value=name)
        ws.cell(row=i, column=5, value=owner)
        ws.cell(row=i, column=7, value=gyoshu)
        ws.cell(row=i, column=9, value=gyomu)
        ws.cell(row=i, column=10, value=nin)
        ws.cell(row=i, column=44, value=mail)
        ws.cell(row=i, column=46, value=num)
        ws.cell(row=i, column=47, value=spec)
    path = DEMO / "demo_perfile_36kyotei.xlsx"
    wb.save(path)
    return path


def build_samples():
    SAMPLES.mkdir(parents=True, exist_ok=True)
    made = []
    for title in ["建設業見本", "製造業見本", "サービス業見本"]:
        p = SAMPLES / f"{title}.pdf"
        p.write_bytes(_minimal_pdf(title))
        made.append(p)
    return made


if __name__ == "__main__":
    t = build_template()
    d = build_demo()
    s = build_samples()
    print(f"テンプレート: {t}")
    print(f"デモExcel  : {d}")
    print("見本サンプル :")
    for p in s:
        print(f"  - {p}")
    print("生成完了")
